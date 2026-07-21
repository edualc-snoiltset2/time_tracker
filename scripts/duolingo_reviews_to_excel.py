#!/usr/bin/env python3
"""Scrape user reviews from the Google Play Store and export them to Excel.

Defaults to Duolingo (``com.duolingo``) but works for any app id. Reviews are
fetched with the ``google-play-scraper`` library (which talks to the Play
Store's own batched review endpoint) and written to a formatted ``.xlsx`` file
ready for analysis.

Usage
-----
    # Scrape up to 500 Duolingo reviews (the default) into duolingo_reviews.xlsx
    python duolingo_reviews_to_excel.py

    # Customize count, sort order, language/country, and output path
    python duolingo_reviews_to_excel.py --max-reviews 1000 \\
        --sort newest --lang en --country us --out reviews.xlsx

    # Preview the Excel format without any network access (uses sample data)
    python duolingo_reviews_to_excel.py --demo --out sample.xlsx

Dependencies: ``google-play-scraper`` and ``openpyxl``
(see requirements-scraper.txt).

Responsible use: reviews are public but contain user-authored content and
display names. Respect Google's Terms of Service, keep request volumes modest,
and only use the data for legitimate analysis.
"""

from __future__ import annotations

import argparse
import logging
import sys
from datetime import datetime
from typing import Any, Dict, List, Optional

logger = logging.getLogger("duolingo_reviews")

DEFAULT_APP_ID = "com.duolingo"
DEFAULT_MAX_REVIEWS = 500

# Columns written to the sheet, in order: (header, dict key).
COLUMNS = [
    ("Review ID", "reviewId"),
    ("User Name", "userName"),
    ("Rating", "score"),
    ("Review", "content"),
    ("Thumbs Up", "thumbsUpCount"),
    ("App Version", "reviewCreatedVersion"),
    ("Date", "at"),
    ("Developer Reply", "replyContent"),
    ("Reply Date", "repliedAt"),
]


def fetch_reviews(
    app_id: str,
    max_reviews: int,
    sort: str = "newest",
    lang: str = "en",
    country: str = "us",
) -> List[Dict[str, Any]]:
    """Fetch up to ``max_reviews`` reviews for ``app_id`` from the Play Store.

    Handles pagination via the continuation token so counts above the ~200
    per-request limit are supported.
    """
    from google_play_scraper import Sort, reviews

    sort_map = {
        "newest": Sort.NEWEST,
        "rating": Sort.RATING,
        "relevance": Sort.MOST_RELEVANT,
    }
    sort_enum = sort_map.get(sort, Sort.NEWEST)

    collected: List[Dict[str, Any]] = []
    token: Optional[Any] = None

    while len(collected) < max_reviews:
        batch_size = min(200, max_reviews - len(collected))
        result, token = reviews(
            app_id,
            lang=lang,
            country=country,
            sort=sort_enum,
            count=batch_size,
            continuation_token=token,
        )
        if not result:
            logger.info("No more reviews available; stopping early.")
            break

        collected.extend(result)
        logger.info("Fetched %d / %d reviews", len(collected), max_reviews)

        if token is None:
            logger.info("Reached the end of available reviews.")
            break

    return collected[:max_reviews]


def _sample_reviews(count: int) -> List[Dict[str, Any]]:
    """Generate deterministic sample rows for offline/demo previews."""
    base = [
        ("LangLearner22", 5, "Best app for daily practice. The streak keeps me going!"),
        ("BusyParent", 4, "Great content but the ads got more aggressive lately."),
        ("Skeptic99", 2, "Too many notifications and the owl guilt-trips me."),
        ("PolyglotPro", 5, "Widget and stories are excellent additions."),
        ("CasualUser", 3, "Fun, but hearts system slows learning without Super."),
    ]
    rows: List[Dict[str, Any]] = []
    for i in range(count):
        name, score, content = base[i % len(base)]
        rows.append(
            {
                "reviewId": f"sample-{i:04d}",
                "userName": name,
                "score": score,
                "content": content,
                "thumbsUpCount": (i * 7) % 50,
                "reviewCreatedVersion": "5.180.0",
                "at": datetime(2026, 7, (i % 28) + 1, 9, 30),
                "replyContent": ("Thanks for the feedback!" if score <= 2 else None),
                "repliedAt": (datetime(2026, 7, (i % 28) + 1, 12, 0) if score <= 2 else None),
            }
        )
    return rows


def write_excel(
    rows: List[Dict[str, Any]],
    path: str,
    app_id: str,
    sheet_title: str = "Reviews",
) -> None:
    """Write reviews to a formatted .xlsx workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4B3F72")
    wrap = Alignment(vertical="top", wrap_text=True)

    # Header row.
    for col_idx, (header, _key) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    # Data rows.
    for row_idx, review in enumerate(rows, start=2):
        for col_idx, (_header, key) in enumerate(COLUMNS, start=1):
            value = review.get(key)
            if isinstance(value, datetime):
                value = value.strftime("%Y-%m-%d %H:%M")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap

    # Column widths.
    widths = {
        "Review ID": 16,
        "User Name": 20,
        "Rating": 8,
        "Review": 70,
        "Thumbs Up": 11,
        "App Version": 13,
        "Date": 18,
        "Developer Reply": 40,
        "Reply Date": 18,
    }
    for col_idx, (header, _key) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = widths.get(header, 18)

    # Freeze header and enable autofilter.
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"

    # Summary sheet.
    _add_summary_sheet(wb, rows, app_id)

    wb.save(path)
    logger.info("Wrote %d review(s) to %s", len(rows), path)


def _add_summary_sheet(wb: Any, rows: List[Dict[str, Any]], app_id: str) -> None:
    """Add a small summary sheet with rating distribution and averages."""
    from openpyxl.styles import Font

    ws = wb.create_sheet("Summary")
    ws["A1"] = "Google Play Reviews - Summary"
    ws["A1"].font = Font(bold=True, size=14)

    scores = [r.get("score") for r in rows if isinstance(r.get("score"), (int, float))]
    total = len(rows)
    avg = round(sum(scores) / len(scores), 2) if scores else 0

    ws["A3"], ws["B3"] = "App ID", app_id
    ws["A4"], ws["B4"] = "Total reviews", total
    ws["A5"], ws["B5"] = "Average rating", avg

    ws["A7"] = "Rating"
    ws["B7"] = "Count"
    ws["A7"].font = ws["B7"].font = Font(bold=True)
    for star in range(5, 0, -1):
        row = 8 + (5 - star)
        ws.cell(row=row, column=1, value=f"{star} star")
        ws.cell(row=row, column=2, value=sum(1 for s in scores if int(s) == star))

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 40


def parse_args(argv: Optional[List[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Scrape Google Play reviews (default: Duolingo) into Excel.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument("--app-id", default=DEFAULT_APP_ID, help="Play Store app id.")
    parser.add_argument(
        "--max-reviews",
        type=int,
        default=DEFAULT_MAX_REVIEWS,
        help="Maximum number of reviews to collect.",
    )
    parser.add_argument(
        "--sort",
        choices=["newest", "rating", "relevance"],
        default="newest",
        help="Order in which to fetch reviews.",
    )
    parser.add_argument("--lang", default="en", help="Review language code.")
    parser.add_argument("--country", default="us", help="Store country code.")
    parser.add_argument(
        "--out",
        default="duolingo_reviews.xlsx",
        help="Output .xlsx path.",
    )
    parser.add_argument(
        "--demo",
        action="store_true",
        help="Generate sample data without any network access (format preview).",
    )
    parser.add_argument("--verbose", "-v", action="store_true", help="Debug logging.")
    return parser.parse_args(argv)


def main(argv: Optional[List[str]] = None) -> int:
    args = parse_args(argv)
    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
    )

    if args.max_reviews <= 0:
        logger.error("--max-reviews must be a positive integer.")
        return 2

    if args.demo:
        logger.info("Demo mode: generating %d sample review(s).", args.max_reviews)
        rows = _sample_reviews(args.max_reviews)
    else:
        try:
            rows = fetch_reviews(
                app_id=args.app_id,
                max_reviews=args.max_reviews,
                sort=args.sort,
                lang=args.lang,
                country=args.country,
            )
        except ImportError:
            logger.error(
                "google-play-scraper is not installed. Run:\n"
                "    pip install -r requirements-scraper.txt"
            )
            return 1
        except Exception as exc:  # noqa: BLE001 - surface network/API failures
            logger.error("Failed to fetch reviews: %s", exc)
            return 1

    if not rows:
        logger.warning("No reviews collected; nothing to write.")
        return 1

    write_excel(rows, args.out, app_id=args.app_id)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
